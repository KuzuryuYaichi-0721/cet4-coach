from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL
import datetime

GREEN = RGBColor(31, 111, 92)
DARK = RGBColor(35, 53, 47)
MUTED = RGBColor(95, 110, 103)
ORANGE = RGBColor(180, 95, 30)
LIGHT_GREEN = "E7F3EC"
LIGHT_ORANGE = "FFF2CC"
CODE_BG = "F3F0E7"

OUT = str(Path(__file__).resolve().parent.parent / "docs" / "CET4Coach-User-Manual.zh-CN.docx")

_numbered_count = 0
_last_block = "other"

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(5)

for name, size, color in [("Heading 1", 18, GREEN), ("Heading 2", 13, GREEN), ("Heading 3", 11.5, ORANGE)]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 18)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)

def p(text="", style=None, bold=False, color=None, size=None, align=None, space_after=None):
    global _last_block
    _last_block = "other"
    par = doc.add_paragraph(style=style)
    run = par.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par

def h1(text):
    global _last_block
    _last_block = "other"
    return doc.add_heading(text, level=1)

def h2(text):
    global _last_block
    _last_block = "other"
    return doc.add_heading(text, level=2)

def h3(text):
    global _last_block
    _last_block = "other"
    return doc.add_heading(text, level=3)

def bullet(text, level=0):
    global _last_block
    _last_block = "other"
    par = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    run = par.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return par

def numbered(text):
    global _numbered_count, _last_block
    if _last_block != "numbered":
        _numbered_count = 1
    else:
        _numbered_count += 1
    _last_block = "numbered"
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.75)
    par.paragraph_format.first_line_indent = Cm(-0.65)
    par.paragraph_format.space_after = Pt(3)
    run = par.add_run(f"{_numbered_count}. {text}")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return par

def code(text):
    global _last_block
    _last_block = "other"
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.right_indent = Cm(0.3)
    par.paragraph_format.space_before = Pt(3)
    par.paragraph_format.space_after = Pt(5)
    run = par.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(55, 55, 55)
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CODE_BG)
    pPr.append(shd)

def callout(title, body, fill=LIGHT_GREEN):
    global _last_block
    _last_block = "other"
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 100, 140, 100, 140)
    par = cell.paragraphs[0]
    r = par.add_run(title)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.color.rgb = GREEN
    par2 = cell.add_paragraph()
    r2 = par2.add_run(body)
    r2.font.name = "Microsoft YaHei"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(9.5)

def table(headers, rows, widths=None):
    global _last_block
    _last_block = "other"
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        r = par.add_run(h)
        r.bold = True
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(hdr[i], "1F6F5C")
        set_cell_margins(hdr[i])
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            r = par.add_run(str(val))
            r.font.name = "Microsoft YaHei"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            r.font.size = Pt(9)
            set_cell_margins(cells[i])
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        cantSplit = OxmlElement("w:cantSplit")
        trPr.append(cantSplit)
        if row._index == 0:
            tblHeader = OxmlElement("w:tblHeader")
            trPr.append(tblHeader)
    for idx, row in enumerate(tbl.rows):
        if idx < len(tbl.rows) - 1:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True
    if widths:
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return tbl

PAGE_BREAK_LIMIT = 2
_page_break_count = 0

def page_break():
    global _page_break_count
    _page_break_count += 1
    if _page_break_count <= PAGE_BREAK_LIMIT:
        doc.add_page_break()

def add_footer_page_number(paragraph):
    run = paragraph.add_run("CET4Coach 详细使用手册  ·  ")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    fldChar = OxmlElement("w:fldChar")
    fldChar.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()
    r._r.append(fldChar)
    r._r.append(instrText)
    r._r.append(fldChar2)
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

# Footer
footer = section.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_footer_page_number(fp)

# ---------------- cover ----------------
for _ in range(2):
    doc.add_paragraph()
p("CET4Coach", bold=True, color=GREEN, size=22, align=WD_ALIGN_PARAGRAPH.CENTER)
p("大肥鱼版 · 详细使用手册", bold=True, color=DARK, size=28, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
p("1500 高频词 · 20 题限时 · 听力练习 · 错题画像 · 可选 DeepSeek 讲解", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
p("Windows 10 / 11  ·  本地运行  ·  PowerShell 5.1 / 7", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
p("版本 1.2-fish  ·  2026 年 10 月", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
p("项目地址：https://github.com/KuzuryuYaichi-0721/cet4-coach", size=9, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER)
p("下载地址：https://github.com/KuzuryuYaichi-0721/cet4-coach/releases/latest", size=9, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
callout("先读这一句",
        "这份手册默认你使用的是桌面上的 CET4Coach 文件夹。双击 RUN.bat 就能开始；没有 API Key 也能学单词和刷题，只是 AI 讲解会降级为本地判断。",
        LIGHT_ORANGE)
p("重要：代码是 MIT 开源；题库、词典、翻译和真题句等数据不是 MIT 授权，仅供个人学习。详见仓库 DATA_LICENSE.md 与 THIRD_PARTY_NOTICES.md。", size=8.5, color=MUTED)
page_break()

# ---------------- contents ----------------
h1("目录")
for item in [
    "1. 先认识这个工具",
    "2. 桌面文件与启动",
    "3. 第一次运行必须做的事",
    "4. 首页与五个核心入口",
    "5. 单词 30 分钟：一步一步操作",
    "6. 20 题 30 分钟：一步一步操作",
    "7. 错题本、讲解和技巧",
    "8. 打卡月历与每周错题画像",
    "9. 听力素材整理与临时文件",
    "10. 题库导入与缺口报告",
    "11. 设置、诊断和数据备份",
    "12. 90 天学习路线",
    "13. 常见问题排查",
    "14. 隐私、成本和数据边界",
    "15. 附录：命令与文件速查",
]:
    p(item, size=10.5, space_after=3)
doc.add_paragraph()
callout("建议阅读顺序",
        "今天先读第 2—5 章，完成第一次单词任务；明天再读第 6 章开始刷题。第 9—11 章可以等需要整理听力或导入题库时再看。",
        LIGHT_GREEN)
page_break()

# ---------------- 1 ----------------
h1("1. 先认识这个工具")
p("CET4Coach 是一个 Windows 本地四级备考工具，不是网课平台，也不会把整套题库上传到云端。它把每天的学习拆成两个动作：")
bullet("30 分钟单词：先认意思，再做核心词的中译英和造句。")
bullet("30 分钟刷题：20 道选择题，先选答案，再用一句话说理由。")
p("系统会自动记录：已学单词、错词、错题、正确率、连续打卡天数。你不需要自己做 Excel。")
h2("它适合谁")
bullet("英语基础较弱，看到长句就慌，不知道四级先练什么。")
bullet("每天能拿出 30—60 分钟，但不想研究复杂学习软件。")
bullet("希望题目、音频、学习记录都留在本机。")
h2("它不做什么")
bullet("不承诺保过，不替代完整的考试报名和官方真题。")
bullet("不自动下载版权音频；听力素材需要你自己整理放进 library/source_audio。")
bullet("不伪造题库数量。2020—2025 唯一题量不足 90 天每天新题，系统用 7 天后重出的方式补题位。")
callout("一句话理解",
        "你每天只需要做两件事：先完成 30 分钟单词，再决定要不要做 20 题。其他记录、复习、错题画像交给系统。",
        LIGHT_GREEN)
page_break()

# ---------------- 2 ----------------
h1("2. 桌面文件与启动")
h2("2.1 桌面位置")
p("项目已经放到桌面：")
code(r"C:\Users\erwei\Desktop\CET4Coach")
p("打开文件夹后，优先双击 RUN.bat。不要直接双击 main.ps1，因为 RUN.bat 会自动选择 PowerShell 版本并启用执行策略绕过。")
h2("2.2 文件夹里有什么")
table(["文件 / 目录", "用途"], [
    ["RUN.bat", "启动入口，优先用 pwsh，找不到就用 Windows PowerShell。"],
    ["main.ps1", "程序主入口；正常使用不需要手改。"],
    ["config.json", "路径、模型、超时、预算等配置。不要写入 API Key。"],
    ["scripts/", "所有 PowerShell 模块：单词、刷题、听力、导入、打卡、LLM。"],
    ["data/", "题库、词表、错题本、打卡、缓存、断点。整个目录就是你的学习记录。"],
    ["library/source_audio/", "你放原始听力素材的地方。程序不会删除这里。"],
    ["library/audio/", "整理后的 Section 音频和 index.json。"],
    ["docs/", "使用说明、数据来源、LLM 封装说明、90 天路线。"],
    ["tests/", "离线测试；普通学习不需要运行。"],
    ["tools/", "开发/维护工具，如趣味句预生成。"],
], widths=[4.8, 11.5])
h2("2.3 启动的两种方式")
p("方式一：双击 RUN.bat。")
p("方式二：在文件夹空白处按住 Shift 右键，选择“在此处打开 PowerShell 窗口”，运行：")
code(r"powershell -NoProfile -ExecutionPolicy Bypass -File .\main.ps1")
p("如果 Windows 提示“无法加载文件，因为在此系统上禁止运行脚本”，不用改全局执行策略，继续用 RUN.bat 即可；它已经带了 Bypass。")
p("如果文件是从网上下载后被 Windows 标记，先在文件夹里运行一次：")
code(r"Get-ChildItem -Recurse | Unblock-File")
page_break()

# ---------------- 3 ----------------
h1("3. 第一次运行必须做的事")
h2("3.1 检查环境")
table(["项目", "要求", "没有会怎样"], [
    ["Windows", "Windows 10 / 11", "无法运行 WinForms 界面。"],
    ["PowerShell", "Windows PowerShell 5.1 或 PowerShell 7", "RUN.bat 会优先找 pwsh，再退到 powershell。"],
    ["DeepSeek Key", "可选；环境变量 DEEPSEEK_API_KEY", "仍可学单词/刷题，但 AI 评估、讲解、周画像会跳过或本地降级。"],
    ["ffmpeg", "可选；整理整卷听力时需要", "已分 Section 的文件仍可复制归类，但整卷不能切分。"],
    ["PotPlayer", "可选；播放听力时需要", "听力题会无法播放，可先做阅读/单词。"],
], widths=[3.6, 6.2, 6.5])
h2("3.2 设置 DeepSeek Key（可选）")
p("打开 PowerShell，运行：")
code('setx DEEPSEEK_API_KEY "你的Key"')
p("然后关掉 PowerShell 和程序，重新打开。程序只从环境变量读取 Key，不会写进 config、日志或提示词。")
p("不要这样做：不要把 Key 写进 config.json；不要截图分享带 Key 的窗口；不要把 Key 发到 GitHub。")
p("如果暂时没有 Key：直接跳过，程序会用本地判断继续。单词的英译中、英文拼写和句子都会尽量本地判断；刷题理由质量会显示“未评估”。")
h2("3.3 听力工具（可选）")
p("安装 ffmpeg，并确保在 PowerShell 运行 ffmpeg -version 有输出。需要播放时再安装 PotPlayer；程序会自动找常见安装路径，也可以在 config.json 的 listening.potPlayer 填完整路径。")
h2("3.4 入门微课与 3 题小测")
p("第一次打开会强制进入 5—8 分钟微课，共 5 页：")
numbered("四级整卷分值、听力和阅读合计 70%。")
numbered("听力内部：篇章题 14.2 分，是新闻/长对话的两倍。")
numbered("阅读内部：仔细阅读 14.2 分，选词填空 3.55 分；先做仔细阅读。")
numbered("写作和翻译各 106.5 分，靠框架和少错误稳分。")
numbered("整场时间轴：写作 → 听力 → 仔细阅读 → 翻译 → 匹配 → 选词。")
p("最后 3 道题全对才能进主菜单。答错会回到对应页面再讲一遍，不扣分。小测不是考试，目的是让你知道“分在哪里”。")
callout("第一次不要做太多",
        "完成微课后，今天只点“单词 30 分钟”。不要一上来就刷 20 题；先把“每天坐下来 30 分钟”变成习惯。",
        LIGHT_ORANGE)
page_break()

# ---------------- 4 ----------------
h1("4. 首页与五个核心入口")
p("左侧是固定导航；中间是内容区。正文使用大字号，按钮较大，所有等待都有进度提示。")
h2("4.1 首页状态")
bullet("基础打卡：完成单词 30 分钟后打勾。")
bullet("刷题：完成 20 题后额外打勾。")
bullet("连续天数：按基础打卡计算；连续 7/14/30 天会给一句简短鼓励。")
bullet("未设置 Key 时，首页会出现黄色提示；不影响听课、单词和本地刷题。")
h2("4.2 导航入口")
table(["入口", "什么时候用"], [
    ["单词 30 分钟", "每天优先做。会创建/恢复当天单词任务。"],
    ["20 题刷题", "有余力时做。每次 20 题、30 分钟。"],
    ["打卡月历", "看哪天完成、连续天数、周错题画像。"],
    ["题库导入", "换题库、追加题、查看缺口报告时用。"],
    ["听力整理", "把原始音频按 Section 归类，或按时间点切整卷。"],
    ["入门微课", "重看分值结构；可从设置页再次进入。"],
    ["设置与诊断", "看 API Key 是否存在、ffmpeg/PotPlayer 路径、今日调用次数。"],
], widths=[4.5, 11.8])
h2("4.3 断点与退出")
p("单词和刷题都会自动保存断点。点“保存并退出”或直接关窗口，下一次进入会从当前词/当前题继续。退出时会清掉当天临时听力。")
p("如果今天是第一次打开，并且没完成微课，左侧其他入口会被禁用，避免绕过入门课。")
h2("4.4 窗口、字号和滚动")
p("默认窗口约 1400×940，并会自动限制在当前屏幕工作区内。页面标题 30pt、卡片标题 21pt、正文 15—18pt、注释 12—14pt；标题和正文不再一个大小。")
p("翻页、下一题、提交反馈后会自动回到内容顶部；内容比窗口高时用右侧滚动条。若屏幕较小，窗口会缩小，文字会自动换行而不是互相盖住。")
p("本版不再强制 Per-Monitor DPI；在 150%/200% 缩放下由 Windows 整体缩放窗口，避免“字体放大了、坐标没放大”造成的文字堆叠。\n阅读窗、逐词讲解窗按屏幕工作区自动收缩，顶部按钮和正文分成独立行，正文有边距、自动换行和垂直滚动条，不会压到标题栏。")
h2("4.5 大肥鱼出现在哪里")
p("左侧栏固定显示大肥鱼的 Q 版头像和“今天也一起学吧”；下方是夜景全身立绘。答题正确/错误反馈、加载进度窗会换用不同表情。图片占用独立留白，不覆盖题干、选项和按钮。")
p("配色使用 DeepSeek 经典蓝 #4D6BFE、深蓝 #1B2A6B、浅蓝白 #F0F5FF，辅以珊瑚色和星光金。形象只用你提供的 6 张图裁切，不是项目生成，也不是抓取网络素材；详见 docs/MASCOT.md 和 MASCOT_LICENSE.md。")
page_break()

# ---------------- 5 ----------------
h1("5. 单词 30 分钟：一步一步操作")
h2("5.1 开始前")
p("点左侧“单词 30 分钟”。程序会在倒计时开始前后做三件事：")
numbered("根据你的进度，从 750 个核心词和 750 个认识词中各取约 13 个新词。")
numbered("把昨天/更早到期的错词加入复习，总词量最多 48 个。")
numbered("读取今天刷题计划需要的听力，复制到临时工作区。没有听力也不影响单词。")
p("计时从创建任务时开始，到 30:00 自动结束；不是“学完再计时”。")
h2("5.2 每个词的三步")
h3("第一步：英译中")
p("屏幕显示一个英文句子，目标词高亮。你只需要回答“这句话里这个单词是什么意思”，用中文写，不要求拼写。")
bullet("提交后先本地对照核心释义；明显对上就直接给反馈，不花模型调用。")
bullet("模糊/不确定时才调用 DeepSeek；没有 Key 或额度用完会降级本地提示。")
bullet("反馈页显示：英文原句、中文提示/翻译、核心释义、大肥鱼反馈。")
h3("第二步：中译英（仅核心词）")
p("系统显示中文释义，让你写出对应的英文单词或短语。")
bullet("允许小拼写错误：长度 5 个字母以上容错 1 个字符，9 个字母以上容错 2 个字符。")
bullet("句子/短语里包含目标词也算对。")
bullet("这一步默认不调用大模型，节省成本。")
h3("第三步：造句（仅核心词）")
p("用目标词造一个短句，能看出词义即可。")
bullet("先检查句子是否包含目标词、是否有至少 3 个英文词。")
bullet("满足基本条件时才调用大模型判断“有没有体现词义”。")
bullet("没有 Key 时，只要句子包含目标词且长度够，就按通过处理。")
h2("5.3 单词状态")
table(["状态", "含义", "下一步"], [
    ["已掌握", "意思/英文/句子都过关（认识词只要意思过关）", "移出活跃复习。"],
    ["模糊", "意思基本对，但英文或句子还差一点", "次日优先复习。"],
    ["未掌握", "意思没对上，或主动跳过", "次日优先复习。"],
], widths=[3.2, 7.5, 4.8])
h2("5.4 时间不够、想退出、想跳过")
bullet("30 分钟到点：立即停止出新词，未完成词进明天复习，自动清理临时听力并记基础打卡。")
bullet("中途保存并退出：当前词和步骤会保存；下次从断点继续。")
bullet("“先跳过”：当前词标记未掌握并进复习，不会卡住。")
bullet("提前完成今日词量：剩余时间进入错词回炉，不额外加新词；到点才自动收尾。")
h2("5.5 结束时你会看到")
p("完成/模糊/未掌握数量；临时文件清理结果；提示明天优先复习哪些词。清理日志会写：")
code("已删除临时文件 N 个，源音频目录未变动（临时区与原始素材均只读校验通过）。")
page_break()

# ---------------- 6 ----------------
h1("6. 20 题 30 分钟：一步一步操作")
h2("6.1 选题规则")
p("每次 20 题，限时 30 分钟。程序优先选仔细阅读和听力篇章：")
bullet("至少有 12 题来自仔细阅读 + 听力篇章。")
bullet("至少 6 题来自同一主攻题型。")
bullet("2—3 种题型混合。")
bullet("7 天内做过的题不重复；错题优先重出；年份尽量均匀覆盖 2020—2025。")
bullet("听力题使用原音频；没有原音频时用浏览器英语语音朗读（电脑合成音），不再因为音频库为空就整段跳过听力。")
h2("6.2 界面与倒计时")
p("顶部显示剩余时间、题号、已答数量、重点题数量和主攻题型。30 分钟到点自动结束；已答正常记录，未答标记未完成。")
p("阅读题超过 3 分钟、听力题超过 2 分钟，界面会出现一行轻提醒：先选一个，后面看讲解。")
h2("6.3 单题流程")
numbered("阅读题：左侧显示文章，下面是题干和选项。先读题干圈关键词，再回原文找同义替换。")
numbered("听力题：先点“播放听力”。有原音频时 PotPlayer 打开；没有原音频时打开一个小浏览器窗口，用英语语音朗读。听完点“我已听完，显示选项”，播放次数会被记录。")
numbered("选择答案；4 个选项用单选按钮，多于 4 个（匹配/选词）用下拉框。")
numbered("在理由框写一句中文：我选 B，因为原文 however 后面说……")
numbered("点“提交答案和理由”。系统先本地判对错，再调用 DeepSeek 评估理由质量。")
numbered("卡住时，点“我不会，逐词讲解”（听力题显示“听不懂？逐词讲解”），会打开全文翻译和逐词讲解，不直接公布答案。原文太长时，点“全屏阅读”或双击原文，打开独立阅读窗，可用 A-/A+ 放大、全屏查看。\n看完反馈点“下一题”。不点不会自动跳。")
h2("6.4 反馈页")
p("答对：显示正确答案、评估、大肥鱼的一句鼓励和技巧。答错：显示固定结构的讲解：")
bullet("这道题考什么（一句话）。")
bullet("正确答案为什么对（回原文/听力信号词）。")
bullet("你选的选项为什么是陷阱。")
bullet("下次遇到同类题的动作。")
bullet("技巧库条目：适用场景、动作、依据、置信度。")
p("技巧讲解不会出现“语义场”“图式理论”等术语；依据只用官方大纲、历年真题统计规律、公开教材常见方法或经验性技巧。")
h2("6.5 错题与跳过")
bullet("答错：自动进错题本，记录题干、选项、你的答案、理由、年份套卷、讲解摘要、来源。")
bullet("主动跳过：按未完成进复习队列，不计正确率。")
bullet("同题下次答对且理由过关：标记“已攻克”一次；连续 2 次答对才移出活跃错题本。")
bullet("30 分钟到点：未答题全部进待复习队列；完成 20 题才额外记刷题打卡。")
h2("6.6 听力播放注意")
p("有原音频时只从当天临时工作区播放；没有原音频时用浏览器 speechSynthesis 朗读。浏览器语音页有“开始播放”按钮；长听力按句分段排队，不会说两句就停。若自动播放被拦截，点一下“开始播放”。浏览器语音打不开时，可以点“逐词讲解”看听力原文和翻译。只关闭本系统启动的 PotPlayer 进程，不会强制关闭你自己的播放器。")
h2("6.7 第二次做错题：选项打乱与跳过重复")
p("同一道单选题第一次答错后，下一次再出现时，A/B/C/D 的顺序会自动打乱，正确答案也会映射到新的字母。页面会明确提示“这次选项已打乱”，并在题干上方显示：上次你选了什么、上次错误类型是什么。")
p("不要背上次的字母。先看题干问什么，再回原文/录音定位，最后按内容选。选项打乱只针对 4 个选项的单选题；匹配题和选词填空的选项是段落/词库标签，保持原顺序。")
p("一道题完成两次学习后，题目页会出现“这题已学习两遍”的勾选框。默认勾选，提交后直接读已保存的讲解要点，不再重复调用深度讲解；想重新生成，取消勾选即可。反馈页也会出现“跳过重复内容”按钮。")
h2("6.8 每轮学习保存到哪里")
p("每轮单词/刷题完成后，完整记录写入 data/history/；文件名带日期和时间，另有 latest-quiz.json / latest-word.json。内容包括题目、选项、你的答案、理由、对错、错误类型、讲解来源。中途退出仍保存在 data/sessions/，下次继续。\n当天完成单词或 20 题后，程序还会自动保存 data/progress/latest.cet4progress.json，并弹一次“今日进度已保存”的提醒；同日备份为 data/progress/backup-YYYYMMDD.cet4progress.json。建议每周把整个 data 文件夹复制到 U 盘或网盘。")
p("下一轮选题会优先重出错题本里的活跃错题，并在题目页强调上次的错误；这就是“第二次学习按上次错误复习”的来源。")
page_break()

# ---------------- 7 ----------------
h1("7. 错题本、讲解和技巧")
h2("7.1 错题本里有什么")
table(["字段", "作用"], [
    ["题型/年份/套卷", "知道这道题来自哪套卷。"],
    ["题干/选项/标准答案", "下次重出时不需要重新找题。"],
    ["你的答案/理由", "复盘时看当时怎么想。"],
    ["错误类型", "定位错、同义替换没认出、蒙对、没找到位置等。"],
    ["讲解摘要/技巧", "下次只看一句话动作，不背整段答案。"],
    ["masteredConsecutive", "连续答对次数；2 次后移出活跃错题。"],
], widths=[4.2, 12.1])
h2("7.2 深度讲解会做什么")
p("错题会调用 DeepSeek 做一次针对本题的诊断，而不是只给一段通用解析。它会：")
bullet("先判断这道题具体考什么，再引用原文/录音的定位句或信号词。")
bullet("根据你选的选项和理由，推断你为什么被它吸引；不确定时会写“从你的选择看，可能……”。")
bullet("逐项分析 A/B/C/D 每个选项：正确项为什么对，干扰项偷换了什么，无关项为什么答非所问。")
bullet("给出下次遇到同类题的一个动作、一条技巧依据，以及 2 个自测追问。")
p("默认模型是 deepseek-chat，最多 2400 输出 token；可在 config.json 的 llm.explanationModel 换成你的账号支持的更深模型。没有 API Key 时会明确显示“本地解析”，不会把本地模板说成联网深度讲解。")
h2("7.3 讲解会缓存")
p("同一道题、同一个错误选项，第二次看讲解会直接读本地缓存，不再花 DeepSeek 调用。换了选项才重新生成。")
p("回答评估不缓存；周复盘按周缓存；趣味句按词/批写入 JSON。")
h2("7.4 怎么看错题最有效")
numbered("先看“这道题考什么”，不要先看正确答案。")
numbered("回原文/听力信号词，找出定位句。")
numbered("比较你选的选项和定位句：是偷换主语、答非所问，还是同义替换没认出。")
numbered("只记一个下次动作，例如“先找 but 后面”。")
numbered("隔天再做一次；如果只是记住答案，不算掌握。")
h2("7.5 什么时候可以跳过重复内容")
p("同一道错题完成两次学习后，系统会记录 learnedTwice。下一次遇到它时，题目页会出现“这题已学习两遍，跳过重复逐项讲解”的勾选框，默认勾选；提交后直接读保存的要点，不再调用深度讲解。反馈页仍有“跳过重复内容”按钮，想重新生成就取消勾选。")
callout("不要抄答案",
        "错题本不是抄写本。每一题只写一句：我为什么会选错，下次先做什么。",
        LIGHT_ORANGE)
page_break()

# ---------------- 8 ----------------
h1("8. 打卡月历与每周错题画像")
h2("8.1 基础打卡和刷题标记")
bullet("基础打卡：完成单词 30 分钟，哪怕没学完全部词。")
bullet("刷题标记：完成 20 题才加；只做一部分不算。")
bullet("月历：绿色表示有学习记录；带“刷”表示当天还完成了 20 题。")
bullet("连续 7/14/30 天会给一句简短鼓励，不重复刷屏。断了就重新从 1 开始，不需要补两天。")
h2("8.2 本周错题画像")
p("每周日首页会自动生成一次“本周错题画像”，也可以在打卡页手动生成。内容是 3—5 句话：")
bullet("错得最多的 2 个题型。")
bullet("1 个共同错误模式。")
bullet("1 个下周可执行的小动作。")
p("生成结果缓存在 data/cache；同周再次打开不会重复调用。没有错题时不调用模型，直接用本地摘要。")
h2("8.3 周复盘怎么用")
p("不要因为画像里说你“粗心”就自责。只挑一个动作，下一周每次做题前默念一遍。比如“先读题干再读文章”“听力等 but 后面”。")
page_break()

# ---------------- 9 ----------------
h1("9. 听力素材整理与临时文件")
h2("9.1 两种素材")
table(["素材类型", "例子", "处理方式"], [
    ["已分 Section", "2023_06_set1_sectionA.mp3", "直接复制并重命名为 SectionA.mp3，不转码。"],
    ["整卷", "2023.06_set1.mp3、2023年6月第1套.mp3", "读取 data/audio_timepoints.json 的起止秒，用 ffmpeg 切分。"],
], widths=[3.5, 7.0, 5.8])
p("支持 mp3 / wav / m4a。已分 Section 优先快速路径，因为不重新编码，速度更快、音质不变。")
h2("9.2 操作步骤")
numbered("把原始素材放进 library/source_audio。")
numbered("打开“听力整理”，确认 ffmpeg 状态。")
numbered("点“开始整理”。程序扫描文件名/目录，识别年份、月份、套号、Section。")
numbered("已分 Section：复制到 library/audio/<套卷>/SectionA.mp3。")
numbered("整卷：只有时间点明确才切；没有就跳过并写缺失报告，绝不猜。")
numbered("看缺失报告，补素材或补时间点后重新整理。")
h2("9.3 时间点模板")
p("打开 data/audio_timepoints.json，按套卷填各 Section 的起止秒：")
code('''{
  "2023-06-set1": {
    "SectionA": [0, 420],
    "SectionB": [420, 900],
    "SectionC": [900, 1500]
  }
}''')
p("单位是秒；结束必须大于开始。没有时间点程序只提示，不猜测。")
h2("9.4 临时工作区")
p("单词任务开始时，程序把今天刷题计划需要的听力复制到：")
code(r"%TEMP%\CET4Coach\listening_<日期>")
p("任务结束（30 分钟到或主动结束）自动删除音频、manifest、中间文件和临时日志；不删除 source_audio、library/audio 或题库。")
p("清理前会计算库和原始素材的文件数+总字节指纹，清理后复核。日志写在 data/logs/cleanup-日期.log。异常退出后，下次启动会清理超过 1 天的临时目录。")
h2("9.5 听力播放")
p("听力题点击播放后调用本机 PotPlayer，参数类似：")
code(r'potplayermini64.exe "<临时路径>" /close')
p("程序记录本次 PID，退出时只结束自己启动的进程。播放次数显示在题目下方，并传给评估参考。")
page_break()

# ---------------- 10 ----------------
h1("10. 题库导入与缺口报告")
h2("10.1 什么时候需要导入")
bullet("第一次使用：桌面版已带 1,668 道题，不需要导入。")
bullet("想追加 2015—2019 或 2026 年的结构化题目。")
bullet("想用自己的 JSON/SQLite 题库。")
h2("10.2 支持格式")
table(["格式", "来源", "说明"], [
    ["SQLite .db", "cet-practice 的 seed/cet6.db", "读取 papers/questions/section_materials 实际字段。"],
    ["JSON", "同字段导出", "支持 questions 数组或单题数组；字段名按实际字段识别。"],
], widths=[3.5, 6.0, 6.8])
p("导入按“题干+选项”SHA256 去重；不重复导入；旧 questions.json 会备份到 data/backup。")
h2("10.3 缺口报告")
p("2020—2025 年真实数据去重后是 1,668 道唯一客观题；仔细阅读+听力篇章 620 道。90 天需要 1,800 个题位、1,080 个重点题位。")
p("系统不伪造题目，而是：")
bullet("每轮仍保证 12+ 重点题、6+ 同一主攻、2—3 种题型。")
bullet("7 天内不重复；题位不足时允许 7 天前做过的题重出。")
bullet("把缺口写进 data/gap_report.json 和 data/import_report.json。")
p("catteacher0515/cet4-download 主要是 PDF/音频，不是带答案的结构化 JSON；自动 PDF 抽题会把答案错配风险带进学习数据，所以需要人工确认后再导入。")
callout("关于“刷完题库”",
        "题库不足以支持 90 天每天全新题，这不是 bug。你的目标不是刷完所有题，而是把同一类题的动作练熟。",
        LIGHT_ORANGE)
page_break()

# ---------------- 11 ----------------
h1("11. 设置、诊断和数据备份")
h2("11.1 设置页能看什么")
bullet("API Key 是否存在（只显示存在，不显示内容）。")
bullet("模型、base URL、今日调用次数和预算。")
bullet("题库唯一题量、听力索引数量。")
bullet("ffmpeg、PotPlayer 路径。")
bullet("缓存文件数量。")
p("“测试 DeepSeek 连接”会发一条最小请求；没有 Key 会直接返回提示，不发请求。")
h2("11.2 数据文件")
table(["文件", "作用", "能否删"], [
    ["data/questions.json", "题库", "删除后需重新导入。"],
    ["data/words.json", "1,500 词表", "不建议删。"],
    ["data/word_sentences.json", "真题句/趣味句", "不建议删。"],
    ["data/word_progress.json", "已学词、复习队列", "删了会从头开始。"],
    ["data/wrongbook.json", "错题本", "删了会丢错题。"],
    ["data/checkin.json", "打卡记录", "删了会丢连续天数。"],
    ["data/sessions/", "断点", "可删；会从当天开头重来。"],
    ["data/history/", "每轮完整学习档案", "删了会丢历史，但错题本仍在。"],
    ["data/cache/", "DeepSeek 缓存/周画像", "可删，不影响正确性。"],
    ["data/logs/", "日志", "可删。"],
], widths=[4.7, 7.6, 5.0])
h2("11.3 学习档案")
p("每轮单词/20 题完成后，data/history/ 会多一个带时间的文件，并更新 latest-quiz.json / latest-word.json。里面是完整题目、选项、你的答案、理由、对错、错误类型和讲解来源。")
bullet("想看上次为什么错：打开 latest-quiz.json，对比 selected、correctKey、errorType、misconception。")
bullet("想换电脑：整个 data 文件夹一起复制，history 也带走。")
bullet("想轻量备份：至少保留 wrongbook.json、checkin.json、word_progress.json；history 是复盘资料，不是运行必需。")
h2("11.4 备份")
p("最简单：关掉程序，把整个 data 文件夹复制到安全位置。换电脑时连同 library/audio 一起复制。")
p("不要只复制 questions.json；单词进度、错题、打卡都在其他 JSON 里。")
h2("11.5 成本控制")
p("默认每天最多 60 次 DeepSeek 调用。刷题约 20 次评估 + 错题深度讲解；同一题同一选项命中缓存，第二次不再调用。单词约 10—20 次，周复盘 1 次。超过预算会自动降级本地判断，不卡住学习。")
p("一道题学满两次后，默认勾选“跳过重复讲解”，不会再次调用深度模型；想重新生成时取消勾选即可。")
h2("11.6 跨版本学习进度包")
p("设置页最下面有“学习进度包”卡片。程序在每次答题、每轮结束、关闭时自动更新 data/progress/latest.cet4progress.json；点“导出学习进度”可另存一份带日期的 .cet4progress.json。")
p("导入时按“题干 + 选项内容”的 SHA256 指纹匹配，不依赖题库行号。因此：")
bullet("新题库换了 id、只调了选项顺序：按内容指纹仍能找到，进度继续。")
bullet("新版本有、老版本没有的题：导入后保留为进度包题目，并进入老版本题库；对应讲解/原文快照一起带入。")
bullet("单词按 word 合并，打卡按日期取并集；不会因为换题库把已学词、连续天数清零。")
bullet("每次更新都只增加 schemaVersion，不删除旧字段；未来版本继续按字段存在性读取。")
p("换电脑：原电脑导出 .cet4progress.json，新电脑在设置页导入即可。原电脑的 data 文件夹仍建议整体备份。")
page_break()

# ---------------- 12 ----------------
h1("12. 90 天学习路线")
h2("12.1 今天只做 30 分钟")
p("完成微课后，点“单词 30 分钟”。到点自动结束。今天不需要刷 20 题。")
h2("12.2 第一周")
table(["天", "任务", "只要求"], [
    ["D1", "微课 + 单词 30 分钟", "完成，不要求记住多少"],
    ["D2", "单词 + 20 题", "单词做完；不会也先选"],
    ["D3", "单词 + 20 题", "错题看讲解，不抄答案"],
    ["D4", "单词 + 20 题", "注意听力回放次数"],
    ["D5", "单词 + 20 题", "记一个最常错题型"],
    ["D6", "单词 + 20 题", "试着说出为什么错"],
    ["D7", "单词 + 错题回炉", "看周画像；休息也可以"],
], widths=[1.6, 7.0, 8.0])
p("如果某天只有 30 分钟，只做单词。基础打卡靠单词 30 分钟；20 题是额外标记。")
h2("12.3 90 天四段")
table(["阶段", "重点", "目标"], [
    ["第 1—14 天", "习惯 + 基础词", "每天坐下来；知道题型分值。"],
    ["第 15—45 天", "高频词 + 仔细阅读/听力篇章", "定位句找得到，同义替换认得出。"],
    ["第 46—70 天", "新词变少，提速", "20 题尽量 28 分钟内完成；补最弱题型。"],
    ["第 71—90 天", "模拟 + 保分", "只复习错题和技巧，不开新计划。"],
], widths=[3.4, 7.3, 6.6])
h2("12.4 每天 60 分钟的标准动作")
numbered("0—30 分钟：单词，先猜意思，再中译英/造句。")
numbered("30—60 分钟：20 题，先选答案，再用一句话说理由。")
numbered("最后 2 分钟：记一个下次动作；不抄答案。")
p("如果今天很累，只做单词。断了一天不补两天的量，第二天照常打开。")
h2("12.5 考场时间轴（微课里也讲过）")
p("写作 30 分钟 → 听力 25 分钟 → 仔细阅读 25—30 分钟 → 翻译 20—25 分钟 → 长篇匹配 8—10 分钟 → 选词填空 5—7 分钟（不够就猜）。")
page_break()

# ---------------- 13 ----------------
h1("13. 常见问题排查")
h2("13.1 双击没反应")
bullet("确认解压完整，RUN.bat 和 main.ps1 在同一层。")
bullet("右键 RUN.bat → 以管理员身份运行（一般不需要）。")
bullet("打开 PowerShell，进入文件夹，运行 main.ps1；看红字报错。")
bullet("运行 Get-ChildItem -Recurse | Unblock-File 解除“来自网络”的锁定。")
h2("13.2 PowerShell 报“禁止运行脚本”")
p("不要改全局策略。用 RUN.bat，或在当前窗口运行：")
code(r"powershell -NoProfile -ExecutionPolicy Bypass -File .\main.ps1")
h2("13.3 API 报错")
table(["现象", "可能原因", "处理"], [
    ["没有 DEEPSEEK_API_KEY", "环境变量没生效", "setx 后关掉终端/程序重开；或在“设置”看是否存在。"],
    ["401 / 403", "Key 错误、无权限、欠费", "换正确 Key；检查账户状态。"],
    ["429", "请求太频繁/额度限制", "等一会；系统会自动降级本地，不卡学习。"],
    ["超时/连接失败", "网络/DNS/代理", "检查网络；重试；系统有 2 次指数退避重试。"],
], widths=[4.2, 5.4, 7.5])
h2("13.4 听力相关")
bullet("ffmpeg 未找到：加入 PATH，或在 config.json 写完整 ffmpeg/ffprobe 路径。")
bullet("PotPlayer 未找到：在 config.json 的 listening.potPlayer 写 PotPlayerMini64.exe 完整路径。")
bullet("整卷被跳过：缺 data/audio_timepoints.json 的起止秒；补后重试。")
bullet("某题提示缺音频：该 Section 不在 library/audio/index.json；补素材后重新整理。")
h2("13.5 题库/数据相关")
bullet("没有题：确认 data/questions.json 存在；或从题库导入。")
bullet("今天已完成刷题：同一天默认只能完成一轮；想再练可改日期/导入新题，但不建议刷量。")
bullet("想重置全部进度：关程序，备份并删除 data/word_progress.json、wrongbook.json、checkin.json、sessions、plans，再启动。")
bullet("想换电脑：复制整个 CET4Coach 文件夹，尤其是 data 和 library。")
h2("13.6 语言/字体问题")
p("中文界面依赖 Microsoft YaHei；Windows 10/11 自带。如果字体缺失，安装常见中文字体后重启。")
page_break()

# ---------------- 14 ----------------
h1("14. 隐私、成本和数据边界")
h2("14.1 本地保存什么")
p("题库、音频、单词进度、错题本、打卡、会话断点、缓存、日志都在本机。")
h2("14.2 会上传什么")
p("只有调用 DeepSeek 时，发送当前一道题的题干/选项/理由、或当前单词/句子、或本周错题摘要。不上传整库、音频或身份信息。")
p("API Key 只从环境变量读取；不写入 config、缓存、日志或提示词。日志只记录用途、模型、耗时、重试次数、token 估算、是否命中缓存。")
h2("14.3 数据授权")
p("代码 MIT。题库、词典、翻译和真题句不是 MIT 授权；仅供个人学习。不要在商业产品中重新分发这些数据；详见 DATA_LICENSE.md。")
h2("14.4 成本")
p("每天默认最多 60 次调用。缓存和本地判断会减少调用。没有 Key 或预算用完时，学习流程不阻塞。")
page_break()

# ---------------- 15 ----------------
h1("15. 附录：命令与文件速查")
h2("15.1 常用命令")
code('''# 启动
powershell -NoProfile -ExecutionPolicy Bypass -File .\\main.ps1

# 设置 DeepSeek Key（重开终端后生效）
setx DEEPSEEK_API_KEY "你的Key"

# 解除下载文件锁定
Get-ChildItem -Recurse | Unblock-File

# 语法自检
powershell -NoProfile -ExecutionPolicy Bypass -File .\\tests\\Test-Parse.ps1

# 核心数据/选题自检
powershell -NoProfile -ExecutionPolicy Bypass -File .\\tests\\Test-Core.ps1

# 趣味句一次性预生成（约 50 次调用）
powershell -NoProfile -ExecutionPolicy Bypass -File .\\tools\\Generate-WordSentences.ps1 -BatchSize 30''')
h2("15.2 关键路径")
table(["路径", "说明"], [
    ["C:\\Users\\erwei\\Desktop\\CET4Coach", "桌面版项目。"],
    ["data\\sessions\\word-日期.json", "单词断点。"],
    ["data\\sessions\\quiz-日期.json", "刷题断点。"],
    ["data\\plans\\quiz-日期.json", "当天 20 题计划。"],
    ["data\\logs\\llm.jsonl", "脱敏后的 LLM 调用日志。"],
    ["data\\cache\\", "讲解/周画像缓存，可删。"],
    ["%TEMP%\\CET4Coach\\listening_日期\\", "当天临时听力，结束会删。"],
], widths=[9.0, 7.5])
h2("15.3 每日检查清单")
bullet("单词 30 分钟完成。")
bullet("错题只看 1 个共同错误模式。")
bullet("听力题记录回放次数，不无限重听。")
bullet("到点停，不拖到第二天。")
bullet("需要时备份 data 文件夹。")
doc.add_paragraph()
callout("最后一句",
        "你不需要一次变强。每天 30 分钟单词、能做就加 20 题；系统会帮你记录和复习。明天照常打开就行。",
        LIGHT_GREEN)
p("CET4Coach 大肥鱼版 v1.2-fish · 详细使用手册 · 2026-10", size=8.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print(OUT)
